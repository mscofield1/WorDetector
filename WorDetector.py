import tkinter as tk
from tkinter import font as tkfont
from tkinter import messagebox, filedialog, ttk, colorchooser
import threading
import subprocess
import sys
import os
import random
import time
import ctypes

# Function for automatic library installation
def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

# Required libraries check
libs = {
    "deep_translator": "deep-translator", 
    "fpdf": "fpdf", 
    "fitz": "pymupdf", 
    "docx": "python-docx",
    "gtts": "gTTS", 
    "pygame": "pygame", 
    "langdetect": "langdetect" 
}

for mod, pkg in libs.items():
    try: 
        __import__(mod)
    except ImportError: 
        install(pkg)

from deep_translator import GoogleTranslator
from fpdf import FPDF
import fitz  
from docx import Document 
from gtts import gTTS
import pygame
from langdetect import detect 

# --- WINDOWS DPI AWARENESS ---
try:
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    pass

class WorDetector:
    def __init__(self, root):
        self.root = root
        self.root.title("WorDetector - Vocabulary Marathon")
        
        # --- RESPONSIVE SIZING & ADAPTIVE FONTS ---
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        
        # 1366x768 ve altı için font ve boyut ayarı
        if screen_width <= 1366:
            app_width = int(screen_width * 0.95)
            app_height = int(screen_height * 0.85)
            text_size = 16
            list_size = 12
        else:
            app_width = 1250
            app_height = 900
            text_size = 18
            list_size = 14
            
        self.root.geometry(f"{app_width}x{app_height}")
        
        pygame.mixer.init()
        self.is_dark_mode = True 
        self.set_theme_colors()
        self.current_highlight_color = "#fff200" 
        self.word_tag_map = {} 

        self.code_to_lang = {"en": "English", "tr": "Turkish", "de": "German", "fr": "French", "es": "Spanish", "it": "Italian"}
        self.languages = {v: k for k, v in self.code_to_lang.items()}
        self.root.configure(bg=self.bg_main)

        # Fonts
        self.header_font = tkfont.Font(family="Helvetica", size=14, weight="bold")
        self.text_font = tkfont.Font(family="Georgia", size=text_size)
        self.list_font = tkfont.Font(family="Verdana", size=list_size)

        # --- GRID CONFIGURATION (ÇOK KRİTİK: Sağ panelin taşmasını engeller) ---
        self.root.grid_columnconfigure(0, weight=1) # Sol Panel: Kalan tüm alanı doldurur
        self.root.grid_columnconfigure(1, weight=0, minsize=300) # Sağ Panel: Sabit ve görünür kalır
        self.root.grid_rowconfigure(0, weight=1)

        # --- LEFT PANEL ---
        self.left_frame = tk.Frame(root, bg=self.bg_main)
        self.left_frame.grid(row=0, column=0, sticky="nsew", padx=(20, 10), pady=20)
        
        self.top_panel = tk.Frame(self.left_frame, bg=self.bg_main)
        self.top_panel.pack(fill="x", pady=(0, 10))

        self.btn_theme = tk.Button(self.top_panel, text="🌓 Theme", command=self.toggle_theme, bg="#000000", fg="#000000", relief="flat", padx=10)
        self.btn_theme.pack(side="left", padx=(0, 15))

        self.src_lang_box = ttk.Combobox(self.top_panel, values=list(self.languages.keys()), state="readonly", width=12)
        self.src_lang_box.set("English")
        self.src_lang_box.pack(side="left", padx=5)

        self.dest_lang_box = ttk.Combobox(self.top_panel, values=list(self.languages.keys()), state="readonly", width=12)
        self.dest_lang_box.set("Turkish")
        self.dest_lang_box.pack(side="left", padx=5)

        self.color_preview = tk.Button(self.top_panel, text="Marker Color", command=self.choose_color, bg=self.current_highlight_color, width=12, relief="flat")
        self.color_preview.pack(side="left", padx=10)

        self.btn_open_file = tk.Button(self.left_frame, text="📁 OPEN FILE (PDF / WORD / TXT)", command=self.open_file,
                                     bg="#2c3e50", fg="#000000", font=("Arial", 11, "bold"), relief="flat", pady=10)
        self.btn_open_file.pack(fill="x", pady=(0, 10))

        self.text_container = tk.Frame(self.left_frame, bg=self.bg_main)
        self.text_container.pack(fill="both", expand=True)
        
        self.text_area = tk.Text(self.text_container, wrap="word", font=self.text_font, 
                                 bg=self.bg_widgets, fg=self.fg_text, insertbackground="white", 
                                 padx=15, pady=15, relief="flat")
        self.text_area.pack(side="left", fill="both", expand=True)
        
        self.text_scroll = ttk.Scrollbar(self.text_container, command=self.text_area.yview)
        self.text_scroll.pack(side="right", fill="y")
        self.text_area.configure(yscrollcommand=self.text_scroll.set)
        
        self.text_area.bind("<<Modified>>", self.on_text_modified)

        # --- RIGHT PANEL (Layout Fix applied here) ---
        self.right_frame = tk.Frame(root, width=320, bg=self.bg_main) # Genişlik optimize edildi
        self.right_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 20), pady=20)
        self.right_frame.grid_propagate(False)

        tk.Label(self.right_frame, text="VOCABULARY LIST", font=self.header_font, bg=self.bg_main, fg="#ffffff").pack(pady=(0, 10))

        self.list_container = tk.Frame(self.right_frame, bg=self.bg_widgets)
        self.list_container.pack(fill="both", expand=True)

        self.word_listbox = tk.Listbox(self.list_container, font=self.list_font, 
                                       bg=self.bg_widgets, fg="#ffffff", borderwidth=0, highlightthickness=0)
        self.word_listbox.pack(side="left", fill="both", expand=True)
        
        self.list_scroll = ttk.Scrollbar(self.list_container, command=self.word_listbox.yview)
        self.list_scroll.pack(side="right", fill="y")
        self.word_listbox.configure(yscrollcommand=self.list_scroll.set)

        self.word_listbox.bind("<Double-Button-1>", self.delete_selected_word)
        self.word_listbox.bind("<<ListboxSelect>>", self.play_word_sound)

        # --- ACTION PANEL ---
        self.btn_frame = tk.Frame(self.right_frame, bg=self.bg_main)
        self.btn_frame.pack(fill="x", pady=(15, 0))

        actions = [("QUIZ MODE 🧠", self.open_quiz), ("SAVE TXT", self.save_txt), ("SAVE PDF", self.save_pdf), ("CLEAR LIST", self.clear_list)]
        for text, cmd in actions:
            btn_color = "#2980b9" if "QUIZ" in text else "#34495e"
            tk.Button(self.btn_frame, text=text, command=cmd, bg=btn_color, fg="#000000", font=("Arial", 11, "bold"), relief="flat", pady=8).pack(fill="x", pady=2)

        self.text_area.bind("<ButtonRelease-1>", self.on_word_select)

    def set_theme_colors(self):
        if self.is_dark_mode:
            self.bg_main, self.bg_widgets, self.fg_text, self.insert_bg = "#1e1e1e", "#2d2d2d", "#e0e0e0", "white"
        else:
            self.bg_main, self.bg_widgets, self.fg_text, self.insert_bg = "#f0f0f0", "#ffffff", "#2c3e50", "black"

    def toggle_theme(self):
        self.is_dark_mode = not self.is_dark_mode
        self.set_theme_colors()
        self.root.configure(bg=self.bg_main)
        self.text_area.configure(bg=self.bg_widgets, fg=self.fg_text, insertbackground=self.insert_bg)
        self.word_listbox.configure(bg=self.bg_widgets, fg="#ffffff" if self.is_dark_mode else "#000000")
        self.left_frame.configure(bg=self.bg_main)
        self.right_frame.configure(bg=self.bg_main)
        self.top_panel.configure(bg=self.bg_main)

    def on_word_select(self, event):
        try:
            if self.text_area.tag_ranges(tk.SEL):
                start, end = self.text_area.index(tk.SEL_FIRST), self.text_area.index(tk.SEL_LAST)
                word = self.text_area.get(start, end).strip()
                if word and len(word.split()) <= 3:
                    tag_id = f"tag_{int(time.time() * 1000)}"
                    self.text_area.tag_add(tag_id, start, end)
                    self.text_area.tag_configure(tag_id, background=self.current_highlight_color, foreground="black")
                    threading.Thread(target=self.translate_word, args=(word, tag_id), daemon=True).start()
        except: pass

    def translate_word(self, word, tag_id):
        try:
            src, dest = self.languages[self.src_lang_box.get()], self.languages[self.dest_lang_box.get()]
            ans = GoogleTranslator(source=src, target=dest).translate(word)
            pos = ""
            if src == "en":
                low = word.lower()
                if low.endswith(("ing", "ed", "ate", "ize")): pos = " (v)"
                elif low.endswith(("ly")): pos = " (adv)"
                elif low.endswith(("ful", "able", "ous", "ive", "al")): pos = " (adj)"
                else: pos = " (n)"
            list_text = f"➜ {word.lower()}{pos} : {ans.lower()}"
            self.word_listbox.insert(0, list_text)
            self.word_tag_map[list_text] = tag_id
        except: pass

    def delete_selected_word(self, event):
        sel = self.word_listbox.curselection()
        if sel:
            list_text = self.word_listbox.get(sel[0])
            if list_text in self.word_tag_map:
                tag_to_remove = self.word_tag_map[list_text]
                self.text_area.tag_remove(tag_to_remove, "1.0", tk.END)
                del self.word_tag_map[list_text]
            self.word_listbox.delete(sel[0])

    def open_file(self):
        path = filedialog.askopenfilename(filetypes=[("Documents", "*.pdf *.docx *.txt")])
        if not path: return
        try:
            if path.endswith(".pdf"):
                doc = fitz.open(path); text = "".join([p.get_text() for p in doc])
            elif path.endswith(".docx"):
                doc = Document(path); text = "\n".join([p.text for p in doc.paragraphs])
            else:
                with open(path, "r", encoding="utf-8") as f: text = f.read()
            self.text_area.delete("1.0", tk.END); self.text_area.insert("1.0", text)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to open file: {e}")

    def on_text_modified(self, event):
        if self.text_area.edit_modified():
            content = self.text_area.get("1.0", "1.500").strip()
            if len(content) > 10:
                try:
                    lang = detect(content)
                    if lang in self.code_to_lang: self.src_lang_box.set(self.code_to_lang[lang])
                except: pass
            self.text_area.edit_modified(False)

    def play_word_sound(self, event):
        sel = self.word_listbox.curselection()
        if not sel: return
        word = self.word_listbox.get(sel[0]).split(":")[0].replace("➜", "").split("(")[0].strip()
        
        def speak():
            try:
                pygame.mixer.music.stop()
                pygame.mixer.music.unload()
                tts = gTTS(text=word, lang=self.languages[self.src_lang_box.get()])
                tts.save("temp.mp3")
                pygame.mixer.music.load("temp.mp3")
                pygame.mixer.music.play()
            except: pass
        threading.Thread(target=speak, daemon=True).start()

    def open_quiz(self):
        items = list(self.word_listbox.get(0, tk.END))
        if len(items) < 4:
            messagebox.showwarning("Quiz", "You need at least 4 words for a Marathon Quiz!")
            return
        
        remaining_questions = items.copy()
        random.shuffle(remaining_questions)
        total_q = len(remaining_questions)
        score = [0]; errors = []

        quiz_win = tk.Toplevel(self.root)
        quiz_win.title("WorDetector - Vocabulary Marathon")
        quiz_win.geometry("500x650")
        quiz_win.configure(bg=self.bg_widgets)

        self.quiz_ui = {} 
        tk.Label(quiz_win, text="VOCABULARY MARATHON", font=("Arial", 12, "bold"), bg=self.bg_widgets, fg="#2980b9").pack(pady=10)
        self.quiz_ui['stats_label'] = tk.Label(quiz_win, text=f"Question: 1 / {total_q}", bg=self.bg_widgets, fg=self.fg_text)
        self.quiz_ui['stats_label'].pack()
        self.quiz_ui['word_label'] = tk.Label(quiz_win, text="", font=("Georgia", 24, "bold"), bg=self.bg_widgets, fg="#fff200")
        self.quiz_ui['word_label'].pack(pady=15)
        self.quiz_ui['btn_frame'] = tk.Frame(quiz_win, bg=self.bg_widgets)
        self.quiz_ui['btn_frame'].pack(fill="both", expand=True, padx=40)

        def next_question():
            if not remaining_questions:
                report = f"Marathon Finished!\nFinal Score: {score[0]} / {total_q}\n"
                if errors: report += "\nReview these words:\n" + "\n".join(errors)
                messagebox.showinfo("Quiz Over", report); quiz_win.destroy()
                return

            for widget in self.quiz_ui['btn_frame'].winfo_children(): widget.destroy()
            q_item = remaining_questions.pop(0)
            word = q_item.split(":")[0].replace("➜", "").strip()
            correct_ans = q_item.split(":")[1].strip()
            self.quiz_ui['word_label'].config(text=word)
            self.quiz_ui['stats_label'].config(text=f"Question: {total_q - len(remaining_questions)} / {total_q}")

            all_meanings = [i.split(":")[1].strip() for i in items]
            if correct_ans in all_meanings: all_meanings.remove(correct_ans)
            options = random.sample(list(set(all_meanings)), min(3, len(all_meanings))) + [correct_ans]
            random.shuffle(options)

            for opt in options:
                btn = tk.Button(self.quiz_ui['btn_frame'], text=opt, font=("Arial", 11), 
                                bg="#34495e" if self.is_dark_mode else "#ecf0f1",
                                fg="#000000", relief="flat", pady=10, cursor="hand2")
                btn.configure(command=lambda b=btn, s=opt: check_ans(b, s, correct_ans, word))
                btn.pack(fill="x", pady=5)

        def check_ans(clicked_btn, selected, correct, original_word):
            for b in self.quiz_ui['btn_frame'].winfo_children():
                b.config(state="disabled")
                if b['text'] == correct: b.config(bg="#27ae60", fg="#000000") 
            if selected == correct: score[0] += 1
            else:
                clicked_btn.config(bg="#c0392b", fg="#000000") 
                errors.append(f"• {original_word} → {correct}")
            quiz_win.after(1200, next_question) 

        next_question()

    def choose_color(self):
        c = colorchooser.askcolor()[1]
        if c: self.current_highlight_color = c; self.color_preview.config(bg=c)

    def clear_list(self):
        if messagebox.askyesno("Confirm", "Clear everything?"):
            for tag in self.word_tag_map.values(): self.text_area.tag_remove(tag, "1.0", tk.END)
            self.word_listbox.delete(0, tk.END); self.word_tag_map.clear()

    def save_txt(self):
        p = filedialog.asksaveasfilename(defaultextension=".txt")
        if p:
            with open(p, "w", encoding="utf-8") as f:
                for i in self.word_listbox.get(0, tk.END): f.write(i + "\n")

    def save_pdf(self):
        p = filedialog.asksaveasfilename(defaultextension=".pdf")
        if p:
            try:
                pdf = FPDF(); pdf.add_page(); pdf.set_font("Arial", size=12)
                for i in self.word_listbox.get(0, tk.END):
                    clean = i.replace("➜", "-")
                    cmap = {"ğ":"g","Ğ":"G","ı":"i","İ":"I","ö":"o","Ö":"O","ü":"u","Ü":"U","ş":"s","Ş":"S","ç":"c","Ç":"C"}
                    for k, v in cmap.items(): clean = clean.replace(k, v)
                    pdf.cell(200, 10, txt=clean.encode('latin-1', 'ignore').decode('latin-1'), ln=True)
                pdf.output(p)
                messagebox.showinfo("Success", "Vocabulary PDF has been saved.")
            except Exception as e: messagebox.showerror("Error", str(e))

if __name__ == "__main__":
    root = tk.Tk(); app = WorDetector(root); root.mainloop()
